import pathlib
from docx import Document

# Absolute path to the human‑friendly markdown report (created as an artifact)
md_path = pathlib.Path(r"C:\\Users\\hp\\.gemini\\antigravity\\brain\\e7f9bc94-c3ec-483b-b73e-f23c0a24adec\\Humanized_Project_Report_Autonomous_Cloud_Cost_Optimizer.md")

# Destination .docx file – placed in the project folder for easy access
output_path = pathlib.Path(__file__).parent / "Humanized_Project_Report_Autonomous_Cloud_Cost_Optimizer.docx"

# Read markdown lines
lines = md_path.read_text(encoding="utf-8").splitlines()

doc = Document()

for line in lines:
    stripped = line.strip()
    if stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=0)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith("#### "):
        doc.add_heading(stripped[5:], level=3)
    elif stripped == "":
        doc.add_paragraph("")
    else:
        doc.add_paragraph(stripped)

# Save the Word document
output_path = output_path.resolve()
output_path.parent.mkdir(parents=True, exist_ok=True)

doc.save(output_path)
print(f"Word document generated at: {output_path}")
